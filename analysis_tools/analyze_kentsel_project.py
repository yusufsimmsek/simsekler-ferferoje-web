from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from PIL import Image, ImageSequence
from pypdf import PdfReader

try:
    import openpyxl
except Exception:  # pragma: no cover - reported in tooling section
    openpyxl = None


ROOT = Path(r"D:\İlkbahar İnşaat\KENTSEL DÖNÜŞÜM PROJELENDİRME")
OUT = Path(r"C:\Users\user\OneDrive\Belgeler\New project\project_analysis")

VERSION_MAP = {
    "AC1009": "AutoCAD R12",
    "AC1012": "AutoCAD R13",
    "AC1014": "AutoCAD R14",
    "AC1015": "AutoCAD 2000/2000i/2002",
    "AC1018": "AutoCAD 2004/2005/2006",
    "AC1021": "AutoCAD 2007/2008/2009",
    "AC1024": "AutoCAD 2010/2011/2012",
    "AC1027": "AutoCAD 2013/2014/2015/2016/2017",
    "AC1032": "AutoCAD 2018/2019/2020/2021/2022/2023/2024",
}

PDF_TEXT_DIR = OUT / "pdf_text"
PDF_PAGE_DIR = OUT / "pdf_pages"
EXCEL_DIR = OUT / "excel_extracts"
TIFF_DIR = OUT / "tiff_previews"
DWG_STRING_DIR = OUT / "dwg_strings"

Image.MAX_IMAGE_PIXELS = None


def ensure_dirs() -> None:
    for directory in [OUT, PDF_TEXT_DIR, PDF_PAGE_DIR, EXCEL_DIR, TIFF_DIR, DWG_STRING_DIR]:
        directory.mkdir(parents=True, exist_ok=True)


def rel(path: Path) -> str:
    return str(path.relative_to(ROOT))


def slug(text: str, max_len: int = 90) -> str:
    normalized = unicodedata.normalize("NFKD", text)
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    ascii_text = re.sub(r"[^A-Za-z0-9._-]+", "_", ascii_text).strip("._-")
    return (ascii_text or "file")[:max_len]


def stable_name(path: Path, suffix: str = "") -> str:
    digest = hashlib.sha1(str(path).encode("utf-8")).hexdigest()[:10]
    return f"{digest}_{slug(path.stem)}{suffix}"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def json_safe(value: Any) -> Any:
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.hex()
    if isinstance(value, (list, tuple, set)):
        return [json_safe(v) for v in value]
    if isinstance(value, dict):
        return {str(k): json_safe(v) for k, v in value.items()}
    return str(value)


def write_json(path: Path, data: Any) -> None:
    path.write_text(json.dumps(json_safe(data), ensure_ascii=False, indent=2), encoding="utf-8")


def write_text(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8", errors="replace")


def file_inventory() -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for path in sorted(ROOT.rglob("*")):
        if not path.is_file():
            continue
        stat = path.stat()
        rows.append(
            {
                "relative_path": rel(path),
                "absolute_path": str(path),
                "stage": path.relative_to(ROOT).parts[0] if path.relative_to(ROOT).parts else "",
                "directory": rel(path.parent),
                "name": path.name,
                "extension": path.suffix.lower(),
                "size_bytes": stat.st_size,
                "size_mb": round(stat.st_size / (1024 * 1024), 3),
                "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(),
                "sha256": sha256(path),
            }
        )
    return rows


def run_pdftotext(path: Path) -> str | None:
    exe = shutil.which("pdftotext")
    if not exe:
        candidate = Path(r"C:\Program Files\Git\mingw64\bin\pdftotext.exe")
        if candidate.exists():
            exe = str(candidate)
    if not exe:
        return None
    output = PDF_TEXT_DIR / f"{stable_name(path)}.pdftotext.txt"
    try:
        subprocess.run(
            [exe, "-layout", "-enc", "UTF-8", str(path), str(output)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=120,
        )
        return output.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        return f"[pdftotext failed: {exc}]"


def analyze_pdf(path: Path) -> dict[str, Any]:
    item: dict[str, Any] = {"relative_path": rel(path), "name": path.name, "errors": []}
    pypdf_pages: list[dict[str, Any]] = []
    full_parts: list[str] = []
    try:
        reader = PdfReader(str(path), strict=False)
        item["page_count"] = len(reader.pages)
        item["metadata"] = {str(k): str(v) for k, v in (reader.metadata or {}).items()}
        for idx, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                text = ""
                item["errors"].append(f"page {idx} extract_text: {exc}")
            full_parts.append(f"\n\n--- PAGE {idx} ---\n{text}")
            media_box = [float(x) for x in page.mediabox]
            pypdf_pages.append(
                {
                    "page": idx,
                    "text_chars": len(text),
                    "rotation": int(page.get("/Rotate", 0) or 0),
                    "mediabox": media_box,
                    "first_500_chars": text[:500],
                }
            )
        pypdf_text = "".join(full_parts).strip()
        pypdf_path = PDF_TEXT_DIR / f"{stable_name(path)}.pypdf.txt"
        write_text(pypdf_path, pypdf_text)
        item["pypdf_text_file"] = str(pypdf_path)
        item["pypdf_text_chars"] = len(pypdf_text)
        item["pages"] = pypdf_pages
    except Exception as exc:
        item["errors"].append(f"PdfReader: {exc}")
        item["page_count"] = None
        item["metadata"] = {}
        item["pages"] = []

    pdftotext_text = run_pdftotext(path)
    if pdftotext_text is not None:
        item["pdftotext_chars"] = len(pdftotext_text)
        item["pdftotext_file"] = str(PDF_TEXT_DIR / f"{stable_name(path)}.pdftotext.txt")
        item["preferred_text_engine"] = "pdftotext" if len(pdftotext_text) >= item.get("pypdf_text_chars", 0) else "pypdf"
        item["first_1200_chars_preferred"] = (
            pdftotext_text if item["preferred_text_engine"] == "pdftotext" else "".join(full_parts)
        )[:1200]
    else:
        item["pdftotext_chars"] = None
        item["preferred_text_engine"] = "pypdf"
        item["first_1200_chars_preferred"] = "".join(full_parts)[:1200]
    return item


def cell_to_record(cell: Any, data_cell: Any | None = None) -> dict[str, Any]:
    value = cell.value
    cached = data_cell.value if data_cell is not None else None
    record = {
        "cell": cell.coordinate,
        "value": value,
        "cached_value": cached,
        "data_type": cell.data_type,
        "number_format": cell.number_format,
        "style_id": getattr(cell, "style_id", None),
    }
    if isinstance(value, str) and value.startswith("="):
        record["formula"] = value
    return record


def analyze_excel(path: Path) -> dict[str, Any]:
    item: dict[str, Any] = {"relative_path": rel(path), "name": path.name, "errors": []}
    if openpyxl is None:
        item["errors"].append("openpyxl not available")
        return item
    try:
        wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
        wb_data = openpyxl.load_workbook(path, data_only=True, read_only=False)
        item["sheet_names"] = wb.sheetnames
        item["properties"] = {
            "creator": wb.properties.creator,
            "lastModifiedBy": wb.properties.lastModifiedBy,
            "created": wb.properties.created,
            "modified": wb.properties.modified,
            "title": wb.properties.title,
            "subject": wb.properties.subject,
            "description": wb.properties.description,
        }
        sheets: list[dict[str, Any]] = []
        text_lines: list[str] = [f"# {path.name}", ""]
        for ws in wb.worksheets:
            ws_data = wb_data[ws.title]
            non_empty: list[dict[str, Any]] = []
            formulas: list[dict[str, Any]] = []
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    data_cell = ws_data[cell.coordinate]
                    record = cell_to_record(cell, data_cell)
                    non_empty.append(record)
                    if "formula" in record:
                        formulas.append(record)
            sheet_summary = {
                "title": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "dimensions": ws.calculate_dimension(),
                "non_empty_count": len(non_empty),
                "formula_count": len(formulas),
                "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
                "tables": list(ws.tables.keys()),
                "non_empty_cells": non_empty,
                "formulas": formulas,
            }
            sheets.append(sheet_summary)
            text_lines.append(f"## Sheet: {ws.title}")
            text_lines.append(f"Dimension: {ws.calculate_dimension()} | Non-empty: {len(non_empty)} | Formulas: {len(formulas)}")
            for record in non_empty:
                value = record.get("value")
                cached = record.get("cached_value")
                if cached not in (None, value):
                    text_lines.append(f"{record['cell']}: {value!r} (cached={cached!r})")
                else:
                    text_lines.append(f"{record['cell']}: {value!r}")
            text_lines.append("")

            csv_path = EXCEL_DIR / f"{stable_name(path)}__{slug(ws.title)}.csv"
            with csv_path.open("w", newline="", encoding="utf-8-sig") as csvfile:
                writer = csv.writer(csvfile)
                for row in ws.iter_rows(values_only=True):
                    writer.writerow(["" if v is None else v for v in row])

        item["sheets"] = sheets
        json_path = EXCEL_DIR / f"{stable_name(path)}.json"
        txt_path = EXCEL_DIR / f"{stable_name(path)}.txt"
        write_json(json_path, item)
        write_text(txt_path, "\n".join(text_lines))
        item["extract_json"] = str(json_path)
        item["extract_text"] = str(txt_path)
    except Exception as exc:
        item["errors"].append(str(exc))
    return item


def analyze_tiff(path: Path) -> dict[str, Any]:
    item: dict[str, Any] = {"relative_path": rel(path), "name": path.name, "errors": [], "frames": []}
    try:
        with Image.open(path) as im:
            item["format"] = im.format
            item["mode"] = im.mode
            item["info"] = {k: str(v) for k, v in im.info.items()}
            frame_count = getattr(im, "n_frames", 1)
            item["frame_count"] = frame_count
            for idx, frame in enumerate(ImageSequence.Iterator(im), start=1):
                frame_info = {
                    "frame": idx,
                    "size": list(frame.size),
                    "mode": frame.mode,
                    "dpi": frame.info.get("dpi"),
                    "tag_v2": {},
                }
                for tag, value in getattr(frame, "tag_v2", {}).items():
                    tag_name = Image.TiffTags.TAGS_V2.get(tag, str(tag))
                    frame_info["tag_v2"][str(tag_name)] = str(value)[:500]
                preview = frame.copy()
                preview.thumbnail((2400, 2400))
                preview_path = TIFF_DIR / f"{stable_name(path)}__frame_{idx:02d}.png"
                preview.save(preview_path)
                frame_info["preview_png"] = str(preview_path)
                item["frames"].append(frame_info)
    except Exception as exc:
        item["errors"].append(str(exc))
    return item


def printable_ratio(text: str) -> float:
    if not text:
        return 0
    printable = sum(1 for c in text if c.isprintable() and (c.isalnum() or c.isspace() or c in "._-:/\\()[]{}+,;#%&=<>"))
    return printable / len(text)


def unique_preserve(items: list[str], limit: int | None = None) -> list[str]:
    seen: set[str] = set()
    output: list[str] = []
    for item in items:
        cleaned = re.sub(r"\s+", " ", item).strip()
        if len(cleaned) < 4 or cleaned in seen:
            continue
        if printable_ratio(cleaned) < 0.72:
            continue
        seen.add(cleaned)
        output.append(cleaned)
        if limit and len(output) >= limit:
            break
    return output


def extract_strings_from_bytes(data: bytes) -> dict[str, list[str]]:
    ascii_strings = [m.group(0).decode("latin-1", errors="ignore") for m in re.finditer(rb"[\x20-\x7e]{4,}", data)]

    decoded_1254 = data.decode("cp1254", errors="ignore")
    cp1254_strings = re.findall(r"[0-9A-Za-zÇĞİÖŞÜçğıöşü ._:/\\()\[\]{}+,\-;#%&=<>]{4,}", decoded_1254)

    decoded_utf16 = data.decode("utf-16le", errors="ignore")
    utf16_strings = re.findall(r"[0-9A-Za-zÇĞİÖŞÜçğıöşü ._:/\\()\[\]{}+,\-;#%&=<>]{4,}", decoded_utf16)

    return {
        "ascii": unique_preserve(ascii_strings, 10000),
        "cp1254": unique_preserve(cp1254_strings, 10000),
        "utf16le": unique_preserve(utf16_strings, 10000),
    }


def analyze_dwg(path: Path) -> dict[str, Any]:
    item: dict[str, Any] = {"relative_path": rel(path), "name": path.name, "errors": []}
    try:
        data = path.read_bytes()
        header = data[:32]
        version = header[:6].decode("ascii", errors="ignore")
        strings = extract_strings_from_bytes(data)
        all_strings = unique_preserve(strings["ascii"] + strings["cp1254"] + strings["utf16le"], 20000)
        token_counter: Counter[str] = Counter()
        cad_hits: Counter[str] = Counter()
        keywords = [
            "AcDb",
            "Layer",
            "Model",
            "Layout",
            "Block",
            "Text",
            "MText",
            "DIM",
            "HATCH",
            "POLYLINE",
            "LINE",
            "INSERT",
            "A-WALL",
            "AXIS",
            "DUVAR",
            "MERDİVEN",
            "MERDIVEN",
            "ASANSÖR",
            "ASANSOR",
            "KAPI",
            "PENCERE",
            "ADA",
            "PARSEL",
            "PLAN",
            "İMAR",
            "IMAR",
        ]
        for s in all_strings:
            for word in re.findall(r"[A-Za-zÇĞİÖŞÜçğıöşü0-9_./-]{3,}", s):
                token_counter[word.upper()] += 1
            upper = s.upper()
            for keyword in keywords:
                if keyword.upper() in upper:
                    cad_hits[keyword] += 1
        item.update(
            {
                "dwg_header": header.hex(" "),
                "dwg_version_code": version,
                "dwg_version_name": VERSION_MAP.get(version, "Unknown/unsupported header"),
                "ascii_string_count": len(strings["ascii"]),
                "cp1254_string_count": len(strings["cp1254"]),
                "utf16le_string_count": len(strings["utf16le"]),
                "combined_string_count": len(all_strings),
                "top_tokens": token_counter.most_common(80),
                "cad_keyword_hits": cad_hits.most_common(),
                "sample_strings": all_strings[:250],
            }
        )
        string_path = DWG_STRING_DIR / f"{stable_name(path)}.strings.txt"
        write_text(string_path, "\n".join(all_strings))
        item["strings_file"] = str(string_path)
    except Exception as exc:
        item["errors"].append(str(exc))
    return item


def group_inventory(inventory: list[dict[str, Any]]) -> dict[str, Any]:
    by_stage: dict[str, dict[str, Any]] = defaultdict(lambda: {"files": 0, "size_bytes": 0, "extensions": Counter()})
    by_ext: Counter[str] = Counter()
    hashes: dict[str, list[str]] = defaultdict(list)
    for row in inventory:
        stage = row["stage"]
        by_stage[stage]["files"] += 1
        by_stage[stage]["size_bytes"] += row["size_bytes"]
        by_stage[stage]["extensions"][row["extension"]] += 1
        by_ext[row["extension"]] += 1
        hashes[row["sha256"]].append(row["relative_path"])
    return {
        "total_files": len(inventory),
        "total_size_mb": round(sum(r["size_bytes"] for r in inventory) / (1024 * 1024), 3),
        "by_extension": dict(by_ext),
        "by_stage": {
            stage: {
                "files": data["files"],
                "size_mb": round(data["size_bytes"] / (1024 * 1024), 3),
                "extensions": dict(data["extensions"]),
            }
            for stage, data in sorted(by_stage.items())
        },
        "duplicate_hashes": {h: paths for h, paths in hashes.items() if len(paths) > 1},
    }


def write_markdown_report(
    inventory: list[dict[str, Any]],
    pdfs: list[dict[str, Any]],
    excels: list[dict[str, Any]],
    tiffs: list[dict[str, Any]],
    dwgs: list[dict[str, Any]],
    grouped: dict[str, Any],
) -> None:
    lines: list[str] = []
    lines.append("# Kentsel Dönüşüm Proje Klasörü - İlk Teknik Envanter")
    lines.append("")
    lines.append(f"Kaynak klasör: `{ROOT}`")
    lines.append(f"Analiz zamanı: {datetime.now(timezone.utc).astimezone().isoformat()}")
    lines.append("")
    lines.append("## Genel Özet")
    lines.append(f"- Toplam dosya: {grouped['total_files']}")
    lines.append(f"- Toplam boyut: {grouped['total_size_mb']} MB")
    lines.append(f"- Uzantılar: {grouped['by_extension']}")
    lines.append("")
    lines.append("## Aşamalara Göre")
    for stage, data in grouped["by_stage"].items():
        lines.append(f"- {stage}: {data['files']} dosya, {data['size_mb']} MB, {data['extensions']}")
    lines.append("")
    lines.append("## Dosya Listesi")
    for row in inventory:
        lines.append(f"- `{row['relative_path']}` | {row['size_mb']} MB | sha256 `{row['sha256'][:16]}...`")
    lines.append("")
    lines.append("## PDF Dosyaları")
    for p in pdfs:
        lines.append(f"### {p['relative_path']}")
        lines.append(f"- Sayfa: {p.get('page_count')} | pypdf chars: {p.get('pypdf_text_chars')} | pdftotext chars: {p.get('pdftotext_chars')}")
        lines.append(f"- Tercih edilen metin motoru: {p.get('preferred_text_engine')}")
        if p.get("errors"):
            lines.append(f"- Hatalar: {p.get('errors')}")
        snippet = (p.get("first_1200_chars_preferred") or "").replace("\n", " ")[:500]
        lines.append(f"- İlk metin parçası: {snippet!r}")
    lines.append("")
    lines.append("## Excel Dosyaları")
    for x in excels:
        lines.append(f"### {x['relative_path']}")
        if x.get("errors"):
            lines.append(f"- Hatalar: {x.get('errors')}")
        lines.append(f"- Sayfalar: {x.get('sheet_names')}")
        for sheet in x.get("sheets", []):
            lines.append(
                f"- {sheet['title']}: {sheet['dimensions']}, {sheet['non_empty_count']} dolu hücre, {sheet['formula_count']} formül"
            )
    lines.append("")
    lines.append("## TIFF/Görsel Dosyaları")
    for t in tiffs:
        lines.append(f"### {t['relative_path']}")
        if t.get("errors"):
            lines.append(f"- Hatalar: {t.get('errors')}")
        lines.append(f"- Format/mode: {t.get('format')}/{t.get('mode')} | frame: {t.get('frame_count')}")
        for frame in t.get("frames", []):
            lines.append(f"- Frame {frame['frame']}: {frame['size']} | preview: `{frame['preview_png']}`")
    lines.append("")
    lines.append("## DWG Dosyaları")
    for d in dwgs:
        lines.append(f"### {d['relative_path']}")
        if d.get("errors"):
            lines.append(f"- Hatalar: {d.get('errors')}")
        lines.append(f"- DWG sürümü: {d.get('dwg_version_code')} ({d.get('dwg_version_name')})")
        lines.append(f"- String sayıları: ascii={d.get('ascii_string_count')}, cp1254={d.get('cp1254_string_count')}, utf16={d.get('utf16le_string_count')}, birleşik={d.get('combined_string_count')}")
        hits = d.get("cad_keyword_hits", [])[:20]
        lines.append(f"- CAD anahtar kelime yakalamaları: {hits}")
        top = d.get("top_tokens", [])[:20]
        lines.append(f"- En sık tokenlar: {top}")
        lines.append(f"- String dökümü: `{d.get('strings_file')}`")
    lines.append("")
    if grouped.get("duplicate_hashes"):
        lines.append("## Birebir Aynı Dosyalar")
        for digest, paths in grouped["duplicate_hashes"].items():
            lines.append(f"- `{digest}`: {paths}")
    write_text(OUT / "MASTER_TECHNICAL_INVENTORY.md", "\n".join(lines))


def main() -> None:
    ensure_dirs()
    if not ROOT.exists():
        raise SystemExit(f"Root folder not found: {ROOT}")
    inventory = file_inventory()
    pdfs = [analyze_pdf(Path(row["absolute_path"])) for row in inventory if row["extension"] == ".pdf"]
    excels = [analyze_excel(Path(row["absolute_path"])) for row in inventory if row["extension"] == ".xlsx"]
    tiffs = [analyze_tiff(Path(row["absolute_path"])) for row in inventory if row["extension"] in {".tif", ".tiff"}]
    dwgs = [analyze_dwg(Path(row["absolute_path"])) for row in inventory if row["extension"] == ".dwg"]
    grouped = group_inventory(inventory)

    write_json(OUT / "inventory.json", inventory)
    write_json(OUT / "summary.json", grouped)
    write_json(OUT / "pdf_analysis.json", pdfs)
    write_json(OUT / "excel_analysis.json", excels)
    write_json(OUT / "tiff_analysis.json", tiffs)
    write_json(OUT / "dwg_analysis.json", dwgs)
    write_markdown_report(inventory, pdfs, excels, tiffs, dwgs, grouped)
    print(f"Analysis complete: {OUT}")
    print(json.dumps(grouped, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
